import io
import os
import unicodedata
from copy import copy
from collections import defaultdict
from datetime import datetime

import pandas as pd
import openpyxl
import barcode as barcode_lib
from barcode.writer import ImageWriter
from PIL import Image, ImageDraw, ImageFont
from flask import Flask, jsonify, request, render_template, send_file, flash, redirect, url_for
from werkzeug.utils import secure_filename

app = Flask(__name__)
app.secret_key = os.environ.get("SECRET_KEY", "equivalencia-app-secret")

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DATA_DIR = os.environ.get("DATA_DIR") or os.path.join(BASE_DIR, "data")
os.makedirs(DATA_DIR, exist_ok=True)

GRANDE_CACHE_PATH_XLS = os.path.join(DATA_DIR, "planilla_grande.xls")
GRANDE_CACHE_PATH_XLSX = os.path.join(DATA_DIR, "planilla_grande.xlsx")
GRANDE_CACHE_META = os.path.join(DATA_DIR, "planilla_grande.meta")
STOCK_CACHE_PATH = os.path.join(DATA_DIR, "stock.xlsx")
STOCK_CACHE_META = os.path.join(DATA_DIR, "stock.meta")

ALLOWED_EXT = {".xls", ".xlsx"}
FONT_BOLD = os.path.join(BASE_DIR, "fonts", "DejaVuSans-Bold.ttf")
FONT_REGULAR = os.path.join(BASE_DIR, "fonts", "DejaVuSans.ttf")


def _ext(filename):
    return os.path.splitext(filename)[1].lower()


def _encabezado(valor):
    texto = unicodedata.normalize("NFKD", str(valor).strip()).encode("ascii", "ignore").decode()
    return " ".join(texto.upper().split())


def _texto(valor):
    if valor is None or (isinstance(valor, float) and pd.isna(valor)):
        return ""
    try:
        if pd.isna(valor):
            return ""
    except (TypeError, ValueError):
        pass
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    texto = str(valor).strip()
    if texto.lower() in {"nan", "none", "nat", "<na>", "#n/a", "#na"}:
        return ""
    return texto[:-2] if texto.endswith(".0") else texto


_TALLES = {
    "XXS", "XS", "S", "M", "L", "XL", "XXL", "XXXL", "XXXXL",
    "2XL", "3XL", "4XL", "OS", "OSFA", "ALL", "UNICO", "UNICA", "U",
    "SM", "MD", "LG", "2XS", "3XS",
}
for _n in range(0, 55):
    _TALLES.add(str(_n))


def _parece_talle(valor):
    texto = _texto(valor).upper().replace(" ", "")
    if not texto or len(texto) > 6:
        return False
    return texto in _TALLES


def _parece_codigo(valor):
    texto = _texto(valor).replace(" ", "")
    return texto.isdigit() and len(texto) >= 8


def _partir_articulo(valor):
    texto = _texto(valor)
    if not texto:
        return "", ""
    partes = texto.split(maxsplit=1)
    if len(partes) != 2:
        return "", ""
    sku, color = partes[0], partes[1]
    if len(color) > 12 or " " in color.strip():
        return "", ""
    return sku, color


def _celdas(row):
    valores = []
    if hasattr(row, "tolist"):
        valores = list(row.tolist())
    else:
        valores = [row.get(i) for i in range(len(row))]
    return valores


def _detectar_columnas(df):
    encabezados = {
        _encabezado(valor): columna
        for columna, valor in enumerate(df.iloc[0])
        if not pd.isna(valor) and str(valor).strip()
    }
    col_codigo = _buscar_columna(
        encabezados, "CODIGO", "CODIGOS", "CODIGO DE BARRA", "CODIGOS DE BARRA", "BARRA", "BARRAS", "BARCODE"
    )
    col_articulo = _buscar_columna(encabezados, "ARTICULO", "ARTICULOS", "ARTIC", "SKU COLOR", "CONCAT")
    col_sku = _buscar_columna(encabezados, "SKU")
    col_color = _buscar_columna(encabezados, "COLOR")
    cols_talle = [columna for nombre, columna in encabezados.items() if "TALLE" in nombre or nombre in {"TAL", "SIZE"}]
    col_talle = cols_talle[0] if cols_talle else _buscar_columna(encabezados, "TALLE")
    col_descripcion = _buscar_columna(encabezados, "DESCRIPCION", "NOMBRE", "PRODUCTO")
    tiene_encabezados = col_codigo is not None and (
        col_articulo is not None or (col_sku is not None and col_color is not None)
    )
    if tiene_encabezados:
        return {
            "skip_header": True,
            "codigo": col_codigo,
            "articulo": col_articulo,
            "sku": col_sku,
            "color": col_color,
            "talle": col_talle,
            "descripcion": col_descripcion,
        }

    muestra = df.head(80)
    votos_codigo = defaultdict(int)
    votos_articulo = defaultdict(int)
    votos_talle = defaultdict(int)
    votos_desc = defaultdict(int)
    for _, row in muestra.iterrows():
        celdas = _celdas(row)
        for i, valor in enumerate(celdas):
            if _parece_codigo(valor):
                votos_codigo[i] += 1
            sku, color = _partir_articulo(valor)
            if sku and color:
                votos_articulo[i] += 1
            if _parece_talle(valor):
                votos_talle[i] += 1
            texto = _texto(valor)
            if texto and not _parece_codigo(valor) and not _parece_talle(valor) and not _partir_articulo(valor)[0]:
                if len(texto) >= 8:
                    votos_desc[i] += 1

    col_codigo = max(votos_codigo, key=votos_codigo.get) if votos_codigo else 0
    col_articulo = max(votos_articulo, key=votos_articulo.get) if votos_articulo else 1
    col_talle = max(votos_talle, key=votos_talle.get) if votos_talle else None
    col_descripcion = max(votos_desc, key=votos_desc.get) if votos_desc else None
    return {
        "skip_header": False,
        "codigo": col_codigo,
        "articulo": col_articulo,
        "sku": None,
        "color": None,
        "talle": col_talle,
        "descripcion": col_descripcion,
    }


def _talle_en_fila(row, col_talle):
    if col_talle is not None:
        talle = _texto(row.get(col_talle) if hasattr(row, "get") else row.iloc[col_talle] if col_talle < len(row) else "")
        if _parece_talle(talle):
            return talle.upper()
    celdas = _celdas(row)
    encontrados = []
    for i, valor in enumerate(celdas):
        if i == col_talle:
            continue
        if _parece_talle(valor):
            encontrados.append(_texto(valor).upper())
    if not encontrados:
        return ""
    # En las solapas largas el talle suele estar a la derecha.
    return encontrados[-1]


def _sku_color_en_fila(row, cols):
    sku = _texto(row.get(cols["sku"])) if cols["sku"] is not None else ""
    color = _texto(row.get(cols["color"])) if cols["color"] is not None else ""
    if sku and color:
        return sku, color
    if cols["articulo"] is not None:
        sku, color = _partir_articulo(row.get(cols["articulo"]))
        if sku and color:
            return sku, color
    for valor in _celdas(row):
        sku, color = _partir_articulo(valor)
        if sku and color:
            return sku, color
    return "", ""


def _clave_sku(sku):
    return _texto(sku).upper().replace(" ", "")


def _clave_barra(codigo):
    texto = _texto(codigo).replace(" ", "")
    return str(int(texto)) if texto.isdigit() else texto.upper()


def _buscar_columna(headers, *nombres_posibles):
    for nombre in nombres_posibles:
        if nombre in headers:
            return headers[nombre]
    return None


def _read_grande(path):
    """Lee equivalencias en todas las solapas, aunque el talle cambie de columna."""
    ext = _ext(path)
    engine = "xlrd" if ext == ".xls" else "openpyxl"
    xl = pd.ExcelFile(path, engine=engine)

    lookup = defaultdict(list)
    barcode_index = {}
    sku_index = defaultdict(list)
    for sheet in xl.sheet_names:
        df = xl.parse(sheet, header=None, dtype=str)
        if df.empty:
            continue
        cols = _detectar_columnas(df)
        if cols["skip_header"]:
            df = df.iloc[1:]

        for _, row in df.iterrows():
            codigo = _texto(row.get(cols["codigo"]))
            if not _parece_codigo(codigo) and codigo:
                # A veces el código está en otra columna de la misma fila.
                codigo = next(( _texto(v) for v in _celdas(row) if _parece_codigo(v) ), "")
            if not _parece_codigo(codigo):
                continue

            sku, color = _sku_color_en_fila(row, cols)
            if not sku or not color:
                continue

            descripcion = _texto(row.get(cols["descripcion"])) if cols["descripcion"] is not None else ""
            if descripcion.lower() in {"nan", "nam"} or _parece_talle(descripcion):
                descripcion = ""
            talle = _talle_en_fila(row, cols["talle"])
            articulo = f"{sku} {color}"
            lookup[articulo].append((codigo, talle, descripcion))
            producto = {
                "codigo": codigo,
                "sku": sku,
                "color": color,
                "talle": talle,
                "descripcion": descripcion,
            }
            clave = _clave_barra(codigo)
            actual = barcode_index.get(clave)
            if actual is None or (not _parece_talle(actual.get("talle")) and _parece_talle(talle)):
                barcode_index[clave] = producto
            sku_index[_clave_sku(sku)].append(producto)
    return lookup, barcode_index, dict(sku_index)


_LOOKUP_CACHE = {"path": None, "mtime": None, "lookup": None, "barcode": None, "by_sku": None}


def _get_lookup(path):
    global _LOOKUP_CACHE
    mtime = os.path.getmtime(path)
    if _LOOKUP_CACHE["path"] == path and _LOOKUP_CACHE["mtime"] == mtime:
        return _LOOKUP_CACHE["lookup"], _LOOKUP_CACHE["barcode"], _LOOKUP_CACHE["by_sku"]
    lookup, barcode_index, sku_index = _read_grande(path)
    _LOOKUP_CACHE = {
        "path": path,
        "mtime": mtime,
        "lookup": lookup,
        "barcode": barcode_index,
        "by_sku": sku_index,
    }
    return lookup, barcode_index, sku_index


_STOCK_CACHE = {"path": None, "mtime": None, "by_sku": None}


def _get_stock_index(path):
    """Lee SKU, COLOR y TALLE de la planilla de stock (sin precios)."""
    global _STOCK_CACHE
    mtime = os.path.getmtime(path)
    if _STOCK_CACHE["path"] == path and _STOCK_CACHE["mtime"] == mtime:
        return _STOCK_CACHE["by_sku"]

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    hoja = workbook["STOCK"] if "STOCK" in workbook.sheetnames else workbook[workbook.sheetnames[0]]
    filas = hoja.iter_rows(values_only=True)
    try:
        encabezado = next(filas)
    except StopIteration:
        workbook.close()
        raise ValueError("La planilla de stock está vacía.")

    headers = {_encabezado(valor): i for i, valor in enumerate(encabezado) if valor is not None}
    col_sku = _buscar_columna(headers, "SKU", "ARTICULO", "ARTICULOS")
    col_color = _buscar_columna(headers, "COLOR")
    col_talle = _buscar_columna(headers, "TALLE")
    if col_sku is None:
        col_sku = 1
    if col_color is None:
        col_color = 2
    if col_talle is None:
        col_talle = 3

    by_sku = defaultdict(list)
    vistos = set()
    for row in filas:
        if row is None or len(row) <= max(col_sku, col_color, col_talle):
            continue
        sku = _clave_sku(row[col_sku])
        color = _texto(row[col_color])
        talle = _texto(row[col_talle])
        clave = (sku, color.upper(), talle.upper())
        if not sku or clave in vistos:
            continue
        vistos.add(clave)
        by_sku[sku].append({"sku": sku, "color": color, "talle": talle})
    workbook.close()
    _STOCK_CACHE = {"path": path, "mtime": mtime, "by_sku": dict(by_sku)}
    return _STOCK_CACHE["by_sku"]


def _resolver_grande_cache():
    if os.path.exists(GRANDE_CACHE_PATH_XLS):
        return GRANDE_CACHE_PATH_XLS
    if os.path.exists(GRANDE_CACHE_PATH_XLSX):
        return GRANDE_CACHE_PATH_XLSX
    return None


def _resolver_stock_cache():
    return STOCK_CACHE_PATH if os.path.exists(STOCK_CACHE_PATH) else None


def _leer_meta(path):
    if not os.path.exists(path):
        return None
    with open(path, encoding="utf-8") as f:
        return f.read().strip()


def _info_archivo(path, meta_path):
    if path is None:
        return None
    info = _leer_meta(meta_path)
    if info:
        return info
    nombre = os.path.basename(path)
    when = datetime.fromtimestamp(os.path.getmtime(path)).strftime("%d/%m/%Y %H:%M")
    return f"{nombre} - cargada {when}"


def _guardar_meta(path, filename):
    with open(path, "w", encoding="utf-8") as f:
        f.write(f"{secure_filename(filename)} - cargada {datetime.now().strftime('%d/%m/%Y %H:%M')}")


def _dedupe_codigos(codigos):
    vistos = {}
    for c in codigos:
        try:
            clave = int(c)
        except (TypeError, ValueError):
            clave = c
        if clave not in vistos or len(c) < len(vistos[clave]):
            vistos[clave] = c
    return list(vistos.values())


def _buscar_producto(lookup, sku, color, talle):
    if not sku or not color:
        return None, None
    sku_str = _texto(sku)
    color_str = _texto(color)
    talle_str = _texto(talle)
    key = f"{sku_str} {color_str}"
    candidatos = lookup.get(key, [])
    coincidencias = [
        (codigo, desc) for codigo, t, desc in candidatos if t.strip().upper() == talle_str.upper()
    ]
    if not coincidencias:
        return None, None
    codigo = _dedupe_codigos([c for c, _d in coincidencias])[0]
    return codigo, coincidencias[0][1]


def _buscar_codigo_equivalencia(lookup, sku, color, talle):
    """Busca el código por SKU + código de color + talle, sin distinguir mayúsculas."""
    sku_str = _texto(sku)
    color_str = _texto(color)
    talle_str = _texto(talle)
    if not sku_str or not color_str or not talle_str:
        return ""

    clave = f"{sku_str} {color_str}"
    candidatos = lookup.get(clave, [])
    if not candidatos:
        clave_normalizada = (sku_str.upper(), color_str.upper())
        for articulo, opciones in lookup.items():
            sku_eq, color_eq = _partir_articulo(articulo)
            if (sku_eq.upper(), color_eq.upper()) == clave_normalizada:
                candidatos = opciones
                break

    codigos = [
        codigo for codigo, talle_eq, _descripcion in candidatos
        if _texto(talle_eq).upper() == talle_str.upper()
    ]
    codigos = _dedupe_codigos(codigos)
    if len(codigos) == 1:
        return codigos[0]
    if len(codigos) > 1:
        return " / ".join(codigos)
    return "SIN COINCIDENCIA"


def _procesar_stock_patagonia(grande_path, stock_path):
    """Inserta la columna A y completa los códigos en un Excel de stock.

    El archivo se identifica por encabezados, por lo que además del formato Patagonia
    (SKU en H, Cod.Color en I y Talle en C) admite planillas con SKU, COLOR y TALLE.
    """
    lookup, _barcode_index, _by_sku = _get_lookup(grande_path)
    workbook = openpyxl.load_workbook(stock_path)
    hojas_procesadas = 0

    for hoja in workbook.worksheets:
        headers = {
            _encabezado(hoja.cell(row=1, column=columna).value): columna
            for columna in range(1, hoja.max_column + 1)
            if hoja.cell(row=1, column=columna).value is not None
        }
        col_sku = _buscar_columna(headers, "SKU", "CODIGO", "CODIGO ARTICULO")
        col_color = _buscar_columna(headers, "COD.COLOR", "COD COLOR", "CODIGO COLOR", "COLOR")
        col_talle = _buscar_columna(headers, "TALLE", "TAL", "SIZE")
        if not all((col_sku, col_color, col_talle)):
            continue

        # Al insertar A, las columnas que ya detectamos se desplazan una posición.
        hoja.insert_cols(1)
        hojas_procesadas += 1
        encabezado_modelo = hoja.cell(row=1, column=2)
        destino = hoja.cell(row=1, column=1, value="CODIGO DE BARRAS")
        if encabezado_modelo.has_style:
            destino._style = copy(encabezado_modelo._style)
        destino.font = copy(encabezado_modelo.font)
        destino.fill = copy(encabezado_modelo.fill)
        destino.border = copy(encabezado_modelo.border)
        destino.alignment = copy(encabezado_modelo.alignment)
        destino.protection = copy(encabezado_modelo.protection)
        destino.number_format = "@"
        hoja.column_dimensions["A"].width = 18

        for fila in range(2, hoja.max_row + 1):
            codigo = _buscar_codigo_equivalencia(
                lookup,
                hoja.cell(row=fila, column=col_sku + 1).value,
                hoja.cell(row=fila, column=col_color + 1).value,
                hoja.cell(row=fila, column=col_talle + 1).value,
            )
            celda = hoja.cell(row=fila, column=1, value=codigo)
            celda.number_format = "@"

    if not hojas_procesadas:
        raise ValueError(
            "No encontré una hoja con las columnas SKU, código de color y TALLE en la fila 1."
        )

    salida = io.BytesIO()
    workbook.save(salida)
    salida.seek(0)
    return salida


def _generar_imagen_codigo_barra(codigo, module_height_mm=9.0, module_width_mm=0.28):
    codigo = str(codigo).strip()
    solo_digitos = codigo.isdigit()
    if solo_digitos and len(codigo) == 13:
        clase = "ean13"
    elif solo_digitos and len(codigo) == 12:
        clase = "upca"
    elif solo_digitos and len(codigo) == 8:
        clase = "ean8"
    else:
        clase = "code128"

    def _render(tipo):
        BarcodeClass = barcode_lib.get_barcode_class(tipo)
        bc = BarcodeClass(codigo, writer=ImageWriter())
        buf = io.BytesIO()
        bc.write(
            buf,
            options={
                "write_text": False,
                "quiet_zone": 1.0,
                "module_height": module_height_mm,
                "module_width": module_width_mm,
            },
        )
        buf.seek(0)
        return Image.open(buf).convert("L")

    try:
        return _render(clase)
    except Exception:
        return _render("code128")


def _texto_ajustado(draw, texto, font, max_width_px, max_lineas=2):
    palabras = texto.split()
    lineas = []
    actual = ""
    for palabra in palabras:
        prueba = (actual + " " + palabra).strip()
        ancho = draw.textbbox((0, 0), prueba, font=font)[2]
        if ancho <= max_width_px or not actual:
            actual = prueba
        else:
            lineas.append(actual)
            actual = palabra
            if len(lineas) == max_lineas - 1:
                break
    if actual:
        lineas.append(actual)
    lineas = lineas[:max_lineas]
    usado = " ".join(lineas)
    if len(usado) < len(texto):
        ultima = lineas[-1]
        while draw.textbbox((0, 0), ultima + "…", font=font)[2] > max_width_px and len(ultima) > 1:
            ultima = ultima[:-1]
        lineas[-1] = ultima.rstrip() + "…"
    return lineas


def _es_talle_o_duplicado(descripcion, color, talle, sku):
    texto = str(descripcion or "").strip().upper()
    if not texto:
        return True
    if texto in {str(talle or "").strip().upper(), str(color or "").strip().upper(), str(sku or "").strip().upper()}:
        return True
    return texto in {
        "XXS", "XS", "S", "M", "L", "XL", "XXL", "XXXL", "XXXXL",
        "0", "1", "2", "3", "4", "5", "6", "7", "8", "9", "10", "12", "14", "16",
    }


def generar_etiqueta(descripcion, color, talle, codigo, sku="", dpi=300):
    """Genera la etiqueta de 4cm x 2cm como imagen PNG (devuelve BytesIO)."""
    px_mm = dpi / 25.4
    ancho_mm, alto_mm = 40.0, 20.0
    W = round(ancho_mm * px_mm)
    H = round(alto_mm * px_mm)

    img = Image.new("L", (W, H), color=255)
    draw = ImageDraw.Draw(img)

    margen = round(1.3 * px_mm)
    col_izq_ancho = round(21.5 * px_mm)

    font_desc = ImageFont.truetype(FONT_REGULAR, size=round(2.3 * px_mm))
    font_color = ImageFont.truetype(FONT_REGULAR, size=round(2.6 * px_mm))
    font_talle = ImageFont.truetype(FONT_BOLD, size=round(4.6 * px_mm))
    font_sku = ImageFont.truetype(FONT_REGULAR, size=round(2.2 * px_mm))

    y = margen
    max_w = col_izq_ancho - margen

    if not _es_talle_o_duplicado(descripcion, color, talle, sku):
        for linea in _texto_ajustado(draw, str(descripcion).upper(), font_desc, max_w, max_lineas=2):
            draw.text((margen, y), linea, font=font_desc, fill=0)
            y += draw.textbbox((0, 0), linea, font=font_desc)[3] + round(0.6 * px_mm)
        y += round(0.8 * px_mm)

    draw.text((margen, y), str(color).upper(), font=font_color, fill=0)
    y += draw.textbbox((0, 0), str(color).upper(), font=font_color)[3] + round(1.0 * px_mm)

    draw.text((margen, y), str(talle).upper(), font=font_talle, fill=0)
    y += draw.textbbox((0, 0), str(talle).upper(), font=font_talle)[3] + round(0.8 * px_mm)

    if sku:
        draw.text((margen, y), str(sku).strip(), font=font_sku, fill=0)

    codigo_str = str(codigo).strip()
    col_der_x = margen + col_izq_ancho + round(0.6 * px_mm)
    col_der_ancho = W - col_der_x - margen
    espacio_digitos = round(2.6 * px_mm)
    hueco = round(0.35 * px_mm)
    bc_alto_max = max(round(8 * px_mm), H - (2 * margen) - espacio_digitos - hueco)

    bc_img = _generar_imagen_codigo_barra(codigo_str, module_height_mm=7.0, module_width_mm=0.22)
    escala = min(col_der_ancho / bc_img.width, bc_alto_max / bc_img.height)
    bc_w = max(1, round(bc_img.width * escala))
    bc_alto = max(1, round(bc_img.height * escala))
    bc_img = bc_img.resize((bc_w, bc_alto))
    bc_x = col_der_x + max(0, (col_der_ancho - bc_w) // 2)
    img.paste(bc_img, (bc_x, margen))

    font_digitos_size = round(1.7 * px_mm)
    font_digitos = ImageFont.truetype(FONT_REGULAR, size=font_digitos_size)
    while font_digitos_size > 8:
        font_digitos = ImageFont.truetype(FONT_REGULAR, size=font_digitos_size)
        ancho_txt = draw.textbbox((0, 0), codigo_str, font=font_digitos)[2]
        if ancho_txt <= col_der_ancho:
            break
        font_digitos_size -= 1

    bbox = draw.textbbox((0, 0), codigo_str, font=font_digitos)
    txt_h = bbox[3] - bbox[1]
    digitos_y = min(margen + bc_alto + hueco, H - margen - txt_h)
    digitos_x = col_der_x + max(0, (col_der_ancho - (bbox[2] - bbox[0])) // 2)
    draw.text((digitos_x, digitos_y), codigo_str, font=font_digitos, fill=0)

    salida = io.BytesIO()
    img.save(salida, format="PNG", dpi=(dpi, dpi))
    salida.seek(0)
    return salida


def _error_etiqueta(mensaje, status=400):
    if request.headers.get("X-Requested-With") == "fetch" or "application/json" in request.headers.get("Accept", ""):
        return jsonify({"error": mensaje}), status
    flash(mensaje)
    return redirect(url_for("index"))


def _producto_desde_stock(sku, color, talle, lookup, barcode_index, by_sku_grande, by_sku_stock):
    sku = _clave_sku(sku)
    color = _texto(color)
    talle = _texto(talle)

    if color and talle:
        codigo, descripcion = _buscar_producto(lookup, sku, color, talle)
        if not codigo:
            return None, _error_etiqueta(f"No encontré SKU {sku}, color {color} y talle {talle} en equivalencia.")
        return {
            "sku": sku,
            "color": color,
            "talle": talle,
            "codigo": codigo,
            "descripcion": descripcion or "",
        }, None

    opciones = []
    vistos = set()
    candidatos = list(by_sku_grande.get(sku) or [])
    if by_sku_stock and sku in by_sku_stock:
        stock_set = {
            (_texto(op["color"]).upper(), _texto(op["talle"]).upper())
            for op in by_sku_stock[sku]
        }
        filtrados = [
            op for op in candidatos
            if (_texto(op["color"]).upper(), _texto(op["talle"]).upper()) in stock_set
        ]
        if filtrados:
            candidatos = filtrados

    if color:
        candidatos = [op for op in candidatos if _texto(op["color"]).upper() == color.upper()]
    if talle:
        candidatos = [op for op in candidatos if _texto(op["talle"]).upper() == talle.upper()]

    for op in candidatos:
        clave = (op["color"].upper(), op["talle"].upper(), _clave_barra(op["codigo"]))
        if clave in vistos:
            continue
        vistos.add(clave)
        opciones.append({
            "sku": op["sku"],
            "color": op["color"],
            "talle": op["talle"],
            "codigo": op["codigo"],
            "descripcion": op.get("descripcion") or "",
        })

    if not opciones:
        return None, _error_etiqueta(f"No encontré el SKU {sku} en equivalencia" + (" / stock." if by_sku_stock else "."))
    if len(opciones) == 1:
        return opciones[0], None
    return None, jsonify({"opciones": opciones})


def _resolver_producto(datos):
    barcode_scanned = _texto(datos.get("barcode") or "").replace("\r", "").replace("\n", "").replace("\t", "")
    sku = _clave_sku(datos.get("sku") or "")
    color = _texto(datos.get("color") or "")
    talle = _texto(datos.get("talle") or "")

    grande_path = _resolver_grande_cache()
    if not grande_path:
        return None, _error_etiqueta("No hay planilla de equivalencia cargada. Subila primero.")

    try:
        lookup, barcode_index, by_sku_grande = _get_lookup(grande_path)
    except Exception as e:
        return None, _error_etiqueta(f"Error leyendo equivalencia: {e}", 500)

    by_sku_stock = None
    stock_path = _resolver_stock_cache()
    if stock_path:
        try:
            by_sku_stock = _get_stock_index(stock_path)
        except Exception as e:
            return None, _error_etiqueta(f"Error leyendo el stock: {e}", 500)

    if barcode_scanned:
        producto = barcode_index.get(_clave_barra(barcode_scanned))
        if not producto:
            return None, _error_etiqueta(f"No encontré el código {barcode_scanned}. Probá buscarlo por SKU.")
        return producto, None

    if not sku:
        return None, _error_etiqueta("Escaneá un código de barras o ingresá SKU, color y talle.")

    return _producto_desde_stock(sku, color, talle, lookup, barcode_index, by_sku_grande, by_sku_stock)


def _precargar_planillas():
    grande = _resolver_grande_cache()
    if grande:
        try:
            _get_lookup(grande)
        except Exception:
            pass
    stock = _resolver_stock_cache()
    if stock:
        try:
            _get_stock_index(stock)
        except Exception:
            pass


_precargar_planillas()


@app.route("/etiqueta")
def etiqueta():
    return redirect(url_for("index"))


@app.route("/etiqueta/generar", methods=["GET", "POST"])
def etiqueta_generar():
    producto, error = _resolver_producto(request.values)
    if error is not None:
        return error

    try:
        imagen = generar_etiqueta(
            producto.get("descripcion") or "",
            producto["color"],
            producto["talle"],
            producto["codigo"],
            sku=producto["sku"],
        )
    except Exception as e:
        return _error_etiqueta(f"No pude generar la etiqueta: {e}", 500)

    descargar = request.values.get("descargar") == "1"
    nombre = f"etiqueta_{secure_filename(producto['codigo'])}.png"
    return send_file(
        imagen,
        mimetype="image/png",
        as_attachment=descargar,
        download_name=nombre,
    )


@app.route("/", methods=["GET"])
def index():
    grande_path = _resolver_grande_cache()
    stock_path = _resolver_stock_cache()
    return render_template(
        "index.html",
        grande_cached=grande_path is not None,
        grande_info=_info_archivo(grande_path, GRANDE_CACHE_META),
        stock_cached=stock_path is not None,
        stock_info=_info_archivo(stock_path, STOCK_CACHE_META),
    )


@app.route("/subir-equivalencia", methods=["POST"])
def subir_equivalencia():
    grande_file = request.files.get("grande")
    if not grande_file or grande_file.filename == "":
        flash("Subí la planilla de equivalencia.")
        return redirect(url_for("index"))
    if _ext(grande_file.filename) not in ALLOWED_EXT:
        flash("La planilla de equivalencia debe ser .xls o .xlsx.")
        return redirect(url_for("index"))

    grande_ext = _ext(grande_file.filename)
    grande_path = GRANDE_CACHE_PATH_XLS if grande_ext == ".xls" else GRANDE_CACHE_PATH_XLSX
    other_path = GRANDE_CACHE_PATH_XLSX if grande_ext == ".xls" else GRANDE_CACHE_PATH_XLS
    if os.path.exists(other_path):
        os.remove(other_path)
    grande_file.save(grande_path)

    global _LOOKUP_CACHE
    _LOOKUP_CACHE = {"path": None, "mtime": None, "lookup": None, "barcode": None, "by_sku": None}
    try:
        _get_lookup(grande_path)
    except Exception as e:
        flash(f"La planilla se guardó, pero no pude leerla: {e}")
        return redirect(url_for("index"))

    _guardar_meta(GRANDE_CACHE_META, grande_file.filename)
    flash("Planilla de equivalencia cargada correctamente.")
    return redirect(url_for("index"))


@app.route("/procesar-stock", methods=["POST"])
def procesar_stock():
    stock_file = request.files.get("stock")
    if not stock_file or stock_file.filename == "":
        flash("Subí la planilla que querés completar.")
        return redirect(url_for("index"))
    if _ext(stock_file.filename) != ".xlsx":
        flash("La planilla de stock debe ser un archivo .xlsx.")
        return redirect(url_for("index"))

    grande_path = _resolver_grande_cache()
    if not grande_path:
        flash("Primero cargá la planilla de equivalencia.")
        return redirect(url_for("index"))

    stock_bytes = io.BytesIO(stock_file.read())
    stock_bytes.name = stock_file.filename
    try:
        salida = _procesar_stock_patagonia(grande_path, stock_bytes)
    except Exception as e:
        flash(f"No pude procesar la planilla: {e}")
        return redirect(url_for("index"))

    nombre_base = os.path.splitext(secure_filename(stock_file.filename))[0]
    return send_file(
        salida,
        as_attachment=True,
        download_name=f"{nombre_base} - con codigos de barra.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@app.route("/health")
def health():
    return {"status": "ok"}


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=int(os.environ.get("PORT", 8000)), threaded=True)
